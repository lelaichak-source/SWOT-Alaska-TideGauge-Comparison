"""
10_thesis_tables.py
===================
Build the thesis tables as CSV files and one formatted Excel workbook.

    Table 1  Study stations: region, NOAA ID, coordinates, SWOT passes, TWL use
    Table 2  Matched pairs per station and SWOT pass (NTR comparison)
    Table 3  Agreement statistics (n, r, p, bias, RMSE) for NTR, DAC-restored
             NTR and TWL
    Table 4  Regression slopes with 95% CI for NTR and DAC-restored NTR
             (* p < 0.05, ** p < 0.01, *** p < 0.001)

p-values below 0.001 are written as "<0.001".

Outputs (OUTPUT_DIR/tables/):
    table1_stations.csv, table2_match_summary.csv, table3_statistics.csv,
    table4_regression.csv, thesis_tables.xlsx

Run:
    python 10_thesis_tables.py
"""

import numpy as np
import pandas as pd
from scipy import stats

import config
from swot_utils import PIPELINES, load_pairs

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

# Station order used in the thesis tables (grouped by region).
STATION_ORDER = ["Anchorage", "Nikiski", "Seldovia",
                 "Seward", "Valdez", "Cordova",
                 "Kodiak", "Alitak"]


def fmt_p(p):
    """'<0.001' for very small p-values, otherwise p rounded to 3 decimals."""
    if p is None or not np.isfinite(p):
        return ""
    return "<0.001" if p < 0.001 else round(float(p), 3)


def agreement(x, y):
    if len(x) < 3:
        return {}
    r, p = stats.pearsonr(x, y)
    return dict(n=len(x), r=round(r, 3), p=fmt_p(p), bias=round(float(np.mean(y - x)), 3),
                rmse=round(float(np.sqrt(np.mean((y - x) ** 2))), 3))


def slope(x, y):
    n = len(x)
    if n < 3:
        return {}
    lr = stats.linregress(x, y)
    half = stats.t.ppf(0.975, df=n - 2) * lr.stderr
    sig = "***" if lr.pvalue < 0.001 else "**" if lr.pvalue < 0.01 else "*" if lr.pvalue < 0.05 else ""
    return dict(n=n, slope=round(lr.slope, 3), ci_low=round(lr.slope - half, 3),
                ci_high=round(lr.slope + half, 3), r=round(lr.rvalue, 3), p=fmt_p(lr.pvalue),
                sig=sig)


def station_values(df, key, name):
    p = PIPELINES[key]
    sub = df[df["station"] == name]
    return sub[p["x"]].values, sub[p["y"]].values


def table1():
    return pd.DataFrame([{
        "Station": name, "Region": config.STATIONS[name]["region"],
        "NOAA ID": config.STATIONS[name]["id"], "Latitude (N)": config.STATIONS[name]["lat"],
        "Longitude (E)": config.STATIONS[name]["lon"],
        "SWOT Passes": ", ".join(str(p) for p in sorted(config.STATION_PASSES[name])),
        "TWL Analysis": "Yes" if name in config.TWL_STATIONS else "No",
    } for name in STATION_ORDER])


def table2(df_ntr):
    passes = sorted(df_ntr["pass_num"].dropna().unique().astype(int))
    rows = []
    for name in STATION_ORDER:
        sub = df_ntr[df_ntr["station"] == name]
        row = {"Station": name, "Region": config.STATIONS[name]["region"]}
        for p in passes:
            n = int((sub["pass_num"] == p).sum())
            row[f"Pass {p}"] = n if n > 0 else ""
        row["Total pairs"] = len(sub)
        rows.append(row)
    return pd.DataFrame(rows)


def table3(dfs):
    rows = []
    for name in STATION_ORDER:
        row = {"Station": name, "Region": config.STATIONS[name]["region"]}
        for key, label in [("ntr", "NTR"), ("dac", "DAC"), ("twl", "TWL")]:
            s = agreement(*station_values(dfs[key], key, name)) if key in dfs else {}
            row.update({f"{label} n": s.get("n", ""), f"{label} r": s.get("r", ""),
                        f"{label} p": s.get("p", ""), f"{label} Bias (m)": s.get("bias", ""),
                        f"{label} RMSE (m)": s.get("rmse", "")})
        rows.append(row)
    return pd.DataFrame(rows)


def table4(dfs):
    rows = []
    for name in STATION_ORDER:
        row = {"Station": name, "Region": config.STATIONS[name]["region"]}
        for key, label in [("ntr", "NTR"), ("dac", "DAC")]:
            s = slope(*station_values(dfs[key], key, name)) if key in dfs else {}
            row.update({f"{label} n": s.get("n", ""), f"{label} Slope": s.get("slope", ""),
                        f"{label} 95% CI": f"[{s['ci_low']}, {s['ci_high']}]" if s else "",
                        f"{label} r": s.get("r", ""), f"{label} p": s.get("p", ""),
                        f"{label} Sig": s.get("sig", "")})
        rows.append(row)
    return pd.DataFrame(rows)


def write_excel(path, sheets):
    """Write the tables to one workbook with header shading and region colours."""
    if not HAS_OPENPYXL:
        print("[skip] openpyxl not installed - Excel workbook not written (CSVs are).")
        return
    side = Side(style="thin")
    border = Border(left=side, right=side, top=side, bottom=side)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header_fill, alt_fill = PatternFill("solid", fgColor="1F4E79"), PatternFill("solid", fgColor="D6E4F0")
    white = PatternFill("solid", fgColor="FFFFFF")
    region_fill = {"Cook Inlet": PatternFill("solid", fgColor="FCE4E4"),
                   "Prince William Sound": PatternFill("solid", fgColor="E4EEF8"),
                   "Kodiak": PatternFill("solid", fgColor="E4F5EC")}

    wb = openpyxl.Workbook()
    for i, (title, df) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.sheet_view.showGridLines = False
        for c, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=c, value=col)
            cell.fill, cell.font = header_fill, Font(bold=True, color="FFFFFF", size=10)
            cell.alignment, cell.border = center, border
        for r, (_, row) in enumerate(df.iterrows(), start=2):
            for c, col in enumerate(df.columns, start=1):
                val = row[col]
                if isinstance(val, float) and np.isnan(val):
                    val = ""
                cell = ws.cell(row=r, column=c, value=val)
                cell.fill = (region_fill.get(row.get("Region"), white) if col in ("Station", "Region")
                             else alt_fill if r % 2 == 0 else white)
                cell.font, cell.alignment, cell.border = Font(size=10), center, border
        for col in ws.columns:
            width = max((len(str(c.value or "")) for c in col), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 4, 30)
        ws.freeze_panes = "A2"

    # Coloured comparison headers above the Table 3 columns.
    ws = wb["Table 3 - Statistics"]
    ws.insert_rows(1)
    for label, c0, c1, color in [("NTR", 3, 7, "1F78B4"), ("DAC-Restored NTR", 8, 12, "E31A1C"),
                                 ("TWL", 13, 17, "33A02C")]:
        ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c1)
        cell = ws.cell(row=1, column=c0, value=label)
        cell.fill, cell.font = PatternFill("solid", fgColor=color), Font(bold=True, color="FFFFFF", size=11)
        cell.alignment, cell.border = center, border
    wb.save(path)


def main():
    dfs = {}
    for key in ("ntr", "dac", "twl"):
        try:
            dfs[key] = load_pairs(key)
        except FileNotFoundError as e:
            if key == "ntr":
                raise
            print(f"[warn] {e} - {key.upper()} columns will be empty.")

    tables = [("Table 1 - Stations", table1(), "table1_stations.csv"),
              ("Table 2 - Match Summary", table2(dfs["ntr"]), "table2_match_summary.csv"),
              ("Table 3 - Statistics", table3(dfs), "table3_statistics.csv"),
              ("Table 4 - Regression", table4(dfs), "table4_regression.csv")]
    config.TABLE_DIR.mkdir(parents=True, exist_ok=True)
    for _, df, filename in tables:
        df.to_csv(config.TABLE_DIR / filename, index=False)
    write_excel(config.TABLE_DIR / "thesis_tables.xlsx", [(t, df) for t, df, _ in tables])
    print(f"Tables written to {config.TABLE_DIR}")


if __name__ == "__main__":
    main()
